import json
import re
import subprocess
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
import yaml
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .paths import CONFIGS, FIGURES, MANUSCRIPT, RESULTS, ROOT, SUBMISSION

PUBLIC_REPOSITORY_URL = "https://github.com/bougtoir/acdc_boundary_grid"


def _canonicalize_office_archive(path: Path) -> None:
    temporary = path.with_suffix(f"{path.suffix}.canonical")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(
        temporary,
        "w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
    ) as target:
        for name in sorted(source.namelist()):
            entry = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.create_system = 0
            entry.external_attr = 0
            target.writestr(entry, source.read(name))
    temporary.replace(path)


def _move_table_captions_above_tables(document: Document) -> None:
    for paragraph in document.paragraphs:
        if not paragraph.text.startswith("TABLE "):
            continue
        table = paragraph._p.getprevious()
        if table is not None and table.tag == qn("w:tbl"):
            table.addprevious(paragraph._p)


def _set_columns(section, count: int) -> None:
    columns = section._sectPr.xpath("./w:cols")
    cols = columns[0] if columns else OxmlElement("w:cols")
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), "240")
    if not columns:
        section._sectPr.append(cols)


def _equation(paragraph, text: str) -> None:
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    math.append(run)
    math_para.append(math)
    paragraph._p.append(math_para)


def _format_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.67)
    section.bottom_margin = Inches(0.67)
    section.left_margin = Inches(0.67)
    section.right_margin = Inches(0.67)
    _set_columns(section, 2)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(2)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name = "Times New Roman"
    document.styles["Title"].font.size = Pt(18)
    document.styles["Heading 1"].font.size = Pt(10)
    document.styles["Heading 1"].font.bold = True
    document.styles["Heading 2"].font.size = Pt(10)
    document.styles["Heading 2"].font.italic = True
    if "Figure Caption" not in document.styles:
        caption = document.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = "Times New Roman"
        caption.font.size = Pt(8)


def _reference_text(reference: dict, number: int) -> str:
    details = [f'[{number}] {reference["authors"]}, "{reference["title"]},"']
    details.append(f' {reference["journal"]},')
    if reference.get("volume"):
        details.append(f' vol. {reference["volume"]},')
    if reference.get("issue"):
        details.append(f' no. {reference["issue"]},')
    if reference.get("pages"):
        details.append(f' pp. {reference["pages"]},')
    if reference.get("article"):
        details.append(f' Art. no. {reference["article"]},')
    details.append(f' {reference["year"]}, doi: {reference["doi"]}.')
    return "".join(details)


def _headline_values(values: pd.DataFrame) -> dict[str, float | str]:
    return {row.value_id: row.value for row in values.itertuples()}


def _add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def build_manuscript(values: pd.DataFrame, central: pd.DataFrame, validation: pd.DataFrame) -> Path:
    MANUSCRIPT.mkdir(parents=True, exist_ok=True)
    build = MANUSCRIPT / "build"
    build.mkdir(parents=True, exist_ok=True)
    references = yaml.safe_load((CONFIGS / "references.yaml").read_text())["references"]
    value = _headline_values(values)
    document = Document()
    _format_document(document)
    document.core_properties.title = "Attributing AC/DC Boundary Value in Active Low-Voltage Distribution Planning"
    document.core_properties.author = "[AUTHOR NAMES TO BE CONFIRMED]"
    document.core_properties.created = datetime(2026, 9, 24, tzinfo=timezone.utc)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Attributing AC/DC Boundary Value in Active Low-Voltage Distribution Planning")
    authors = document.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    authors.add_run("[AUTHOR NAMES AND IEEE MEMBERSHIP TO BE CONFIRMED]").bold = True
    affiliation = document.add_paragraph()
    affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    affiliation.add_run("[AFFILIATIONS, CITY, COUNTRY, AND CORRESPONDING AUTHOR TO BE CONFIRMED]")
    abstract = document.add_paragraph()
    abstract.add_run("Abstract—").bold = True
    abstract.add_run(
        "Comparisons between inherited AC feeders and jointly redesigned hybrid networks "
        "confound the permission to use DC with conductor and topology changes. This paper "
        "develops a matched architecture-attribution framework in which the AC/DC boundary "
        "is endogenous while all-AC and root-converted DC remain feasible. Three public "
        "SimBench low-voltage feeders are evaluated on original and explicitly synthetic "
        "topologies across four demand/DER portfolios, five native-DC shares, three spatial "
        "allocations, four boundary efficiencies, and five nested design cells. The exact "
        "dynamic program for the declared balanced radial approximation includes line, "
        "endpoint, boundary-efficiency, standby, material, and converter-capacity terms. "
        f"The factorial contains {int(float(value['scenario_cell_count'])):,} rows and "
        f"{int(float(value['joint_cell_count'])):,} joint-optimization cells. Under the "
        f"central setting, joint planning changed modeled annual loss by "
        f"{float(value['mean_hybrid_vs_ac_loss_change_pct']):.2f}% relative to matched "
        "conductor-optimized AC. Feeder-specific discrepancy screening classified "
        f"{int(float(value['discrepancy_feeder_robust_cells']))} joint cells as robustly "
        f"material and {int(float(value['discrepancy_feeder_indeterminate_cells']))} as "
        "indeterminate. Converter ablations show that nominal efficiency alone is "
        "insufficient: standby and endpoint utilization alter boundaries. Nonlinear all-AC "
        "stress checks converged in every case. The results identify conditional planning "
        "mechanisms, not universal DC superiority or deployment-ready economics."
    )
    terms = document.add_paragraph()
    terms.add_run("Index Terms—").bold = True
    terms.add_run(
        "AC/DC distribution, active distribution networks, architecture attribution, "
        "distribution planning, low-voltage DC, power conversion."
    )
    document.add_heading("I. INTRODUCTION", level=1)
    document.add_paragraph(
        "Native-DC photovoltaics, batteries, electronics, and electric-vehicle interfaces "
        "motivate renewed examination of mixed AC/DC distribution. Prior work has optimized "
        "hybrid network configuration, DC-feeder placement, converter siting, expansion, and "
        "temporal operation [1]–[6]. Consequently, assigning AC or DC states to network "
        "elements is not itself a defensible novelty claim."
    )
    document.add_paragraph(
        "A recurring identification problem remains. An inherited AC feeder is frequently "
        "compared with a hybrid design that is also allowed to reroute lines, resize "
        "conductors, or change conversion assumptions. The resulting difference combines "
        "several design freedoms. It cannot be interpreted as the marginal value of DC. "
        "This paper therefore asks a narrower question: after matching topology, conductor "
        "design, service requirements, and converter accounting, when does permission to "
        "create a DC subtree materially change a low-voltage planning solution?"
    )
    document.add_paragraph(
        "The contributions are: 1) a nested estimand that separates conductor resizing, "
        "fixed-design DC permission, joint domain/conductor redesign, and topology "
        "sensitivity; 2) an exact dynamic program for the stated separable approximation, "
        "checked against independent exhaustive enumeration; 3) full-factorial transition "
        "evidence with explicit treatment of structural repetition; 4) discrepancy-based "
        "classification of modeled benefits; 5) component-level converter ablations; and "
        "6) a reproducible evidence chain from archived public inputs to a scalar registry. "
        "The study does not claim universal DC superiority, observed field-optimal topology, "
        "protection adequacy, or deployable absolute economics."
    )
    document.add_paragraph(
        "This scope is intentionally narrower than a universal architecture ranking. The "
        "analysis estimates marginal architecture effects within a transparent screening "
        "model. Material and converter-capacity terms are normalized regularizers, while "
        "electrical loss components remain separately reported. A reader can therefore "
        "distinguish reduced line current, avoided endpoint conversion, conductor resizing, "
        "and added boundary conversion instead of interpreting their sum as a pure DC benefit."
    )
    document.add_paragraph(
        "Fig. 1 summarizes the controlled evidence chain from public feeder inputs through "
        "nested design cells, exact optimization, validation, and registry-backed reporting."
    )
    document.add_picture(str(FIGURES / "figure_0_formulation.png"), width=Inches(3.1))
    caption = document.add_paragraph(style="Figure Caption")
    caption.add_run("Fig. 1. Controlled architecture-attribution workflow.")
    document.add_heading("II. DATA AND METHODS", level=1)
    document.add_heading("A. Public Inputs and Study Cases", level=2)
    document.add_paragraph(
        "The electrical benchmark consists of the rural1, semiurb4, and urban6 low-voltage "
        "scenario-0 networks from SimBench v1.6.1 [7]. The exact wheel is retained in the "
        "local immutable archive and verified by SHA-256. SimBench provides topology, line "
        "parameters, geographic coordinates, loads, and 15-min annual profiles. The public "
        "charging-transaction dataset [8] is used only to check the scale of the high-EV "
        "stress portfolio; invalid or nonpositive-duration records are excluded by a "
        "predeclared rule. The benchmark topology is engineering reference data rather than "
        "observed infrastructure."
    )
    document.add_paragraph(
        "For each feeder, the original radial topology is paired with a synthetic "
        "minimum-geographic-length spanning tree constructed from the benchmark coordinates. "
        "The latter is explicitly a topology-confounding sensitivity, not a reconstruction of "
        "a utility network. Four portfolios represent base demand (K0), high PV (K1), high EV "
        "(K2), and combined PV, EV, and storage peak-shaving (K3). Their annual-energy and "
        "peak multipliers are frozen in the analysis plan."
    )
    document.add_paragraph(
        "Table I lists the frozen factorial dimensions. They are intentionally balanced "
        "scenario levels and are not weighted to represent deployment prevalence."
    )
    _add_table(
        document,
        ["Dimension", "Levels", "Interpretation"],
        [
            ["Feeder", str(int(float(value["config_feeders_count"]))), "Rural, suburban, urban"],
            ["Topology", str(int(float(value["config_topologies_count"]))), "Original, synthetic MST"],
            ["Portfolio", str(int(float(value["config_portfolios_count"]))), "K0–K3 proxies"],
            [
                "Native-DC share",
                str(int(float(value["config_native_dc_shares_count"]))),
                "0%–100% annual energy",
            ],
            [
                "Spatial allocation",
                str(int(float(value["config_spatial_seeds_count"]))),
                "Frozen seed realizations",
            ],
            [
                "Boundary efficiency",
                str(int(float(value["config_converter_efficiencies_count"]))),
                "94%–99% scenarios",
            ],
            [
                "Design cell",
                str(int(float(value["design_case_count"]))),
                "Nested attribution cells",
            ],
        ],
    )
    document.add_paragraph(
        "TABLE I. Frozen factorial dimensions. Levels are design scenarios rather than "
        "estimated population distributions.",
        style="Figure Caption",
    )
    document.add_paragraph(
        "Table II defines the nested comparison cells used to attribute changes without "
        "silently granting the hybrid design additional conductor or topology freedom."
    )
    _add_table(
        document,
        ["Cell", "Topology", "Conductor", "Domain"],
        [
            ["Inherited AC", "Original", "Base", "AC only"],
            ["AC + conductor", "Matched", "Optimized", "AC only"],
            ["Fixed hybrid", "Matched", "Base", "Endogenous"],
            ["Joint hybrid", "Matched", "Optimized", "Endogenous"],
            ["Ideal hybrid", "Matched", "Optimized", "Lossless conversion"],
        ],
    )
    document.add_paragraph(
        "TABLE II. Nested comparison cells used for attribution.", style="Figure Caption"
    )
    document.add_heading("B. Endogenous Boundary Model", level=2)
    document.add_paragraph(
        "Let G=(V,E) be a radial tree rooted at the AC transformer bus r, with every edge "
        "e=(i,j) directed from parent i to child j. Each non-root bus takes domain d_i in "
        "{AC,DC}; edge (i,j) takes d_j. A boundary converter is installed exactly when "
        "d_i differs from d_j. Each edge selects conductor multiplier s_e from the common "
        "finite menu, or s_e=1 in fixed-conductor cells. Frozen downstream annual energy, "
        "peak power, and equivalent loss hours are denoted E_e, P_e, and H_e."
    )
    paragraph = document.add_paragraph()
    _equation(
        paragraph,
        "I_AC,e = 1000 f_P P_e/(√3 V_AC pf),   I_DC,e = 1000 f_P P_e/V_DC",
    )
    document.add_paragraph(
        f"Here f_P is the portfolio peak multiplier, pf is power factor, "
        f"V_AC={float(value['config_ac_voltage_line_line_v']):.0f} V line-to-line, and "
        f"V_DC={float(value['config_dc_voltage_pole_to_pole_v']):.0f} V pole-to-pole. With base "
        "resistance per kilometre rho_e and length l_e, R_e=rho_e l_e/s_e. The coded annual "
        "line losses are 3 I_AC,e I_AC,e R_e H_e/1000 for AC and "
        "2 I_DC,e I_DC,e R_e H_e/1000 for DC."
    )
    paragraph = document.add_paragraph()
    _equation(
        paragraph,
        "E_mis,i = f_E E_i [z_i(1-q_i)+(1-z_i)q_i],   L_end,i = E_mis,i(1/η_end-1)",
    )
    document.add_paragraph(
        "z_i is one for DC, q_i is the native-DC energy share, and f_E is the annual-energy "
        "multiplier. Thus AC-native energy served from DC and DC-native energy served from AC "
        "are treated symmetrically. For a domain transition on edge e, boundary loss is "
        "f_E E_e(1/eta_b-1)+sigma_b f_P P_e 8760 and installed boundary capacity is f_P P_e. "
        f"The frozen endpoint efficiency is {float(value['config_endpoint_efficiency']) * 100:.0f}% "
        f"and standby fraction is {float(value['config_converter_standby_fraction']) * 100:.1f}% "
        "of rated boundary power. These are dimensionless planning assumptions rather than "
        "empirical vendor curves. The model has no chronological direction reversal, "
        "part-load curve, or loss feedback."
    )
    paragraph = document.add_paragraph()
    _equation(
        paragraph,
        "J = Σ_e L_line,e + Σ_i L_end,i + Σ_e L_boundary,e + λ_M Σ_e M_e + λ_C Σ_e C_e",
    )
    document.add_paragraph(
        "M_e=n(d_j)l_e s_e with n(AC)=4 and n(DC)=3. The material and converter-capacity "
        "weights are normalized regularizers, not monetary costs. Electrical loss components "
        "remain separately reported. The conductor-count convention is a controlled scenario "
        "index and does not establish universal grounding or return-path equivalence."
    )
    paragraph = document.add_paragraph()
    _equation(
        paragraph,
        "F_i(d)=L_end,i(d)+Σ_j min_(d',s)[F_j(d')+L_line,ij(d',s)+λ_M M_ij+B_ij(d,d')]",
    )
    document.add_paragraph(
        "F_i(d) is the minimum score of the subtree rooted at i conditional on domain d; "
        "B_ij contains boundary electrical loss and the capacity regularizer when d differs "
        "from d'. Leaves contribute endpoint loss only, the root terminates at F_r(AC), and "
        "stored choices recover domains and conductors. The algorithm evaluates "
        "O(|E||D|^2|S|) transitions and stores O(|V||D|+|E||D|) values. It is exact only "
        "because downstream profiles are fixed and child subtrees are additive."
    )
    paragraph = document.add_paragraph()
    _equation(
        paragraph,
        "u_e = I_e/(I_max,e s_e),   δ_AC,e = I_AC,e R_e/(V_AC/√3),   δ_DC,e = I_DC,e R_e/(V_DC/2)",
    )
    document.add_paragraph(
        "Ampacity utilization and single-edge voltage drop are recorded after optimization. "
        "Configured limits are screening flags, not constraints in the dynamic program. Exact "
        "optimality therefore does not extend to unbalanced voltage feasibility, meshed "
        "operation, reactive power, converter-loss feedback, protection, or reliability."
    )
    document.add_heading("C. Attribution and Validation", level=2)
    document.add_paragraph(
        "For matched loss L, the conductor contrast is (L_AC,opt−L_AC,base)/L_AC,base; the "
        "fixed-design DC-permission contrast is (L_H,base−L_AC,base)/L_AC,base; and the joint "
        "incremental architecture contrast is (L_H,opt−L_AC,opt)/L_AC,opt. A synthetic-topology "
        "sensitivity compares joint loss with the corresponding original-topology result. "
        "These are attribution contrasts and marginal planning changes within the model, not "
        "causal effects or realized utility savings. The electrical-ideal converter case is "
        "reported only as an upper bound."
    )
    document.add_paragraph(
        "Analytical two-bus cases test current and loss equations and endpoint assignment. "
        "Independent exhaustive enumeration on small radial networks tests all permitted "
        "domain and conductor combinations, including all-AC, root-DC, interior-hybrid, "
        "converter-penalty, endpoint-mismatch, and conductor-driven architecture cases. "
        f"Objective agreement is required within a declared numerical tolerance. Across "
        f"{int(float(value['dp_exactness_case_count']))} exactness cases, the largest absolute "
        f"objective difference was {float(value['dp_exactness_max_objective_error']):.2e}. "
        "Architecture agreement is additionally required for unique optima, while tied "
        "solutions are checked for support of the returned assignment."
    )
    document.add_paragraph(
        "For each SimBench feeder, the all-AC approximation is compared with pandapower's "
        "nonlinear balanced AC line loss using nonlinear loss as the denominator. For a joint "
        "cell with modeled loss contrast delta<0 and feeder discrepancy epsilon_f, the primary "
        "screen labels the benefit robustly material when |delta|>epsilon_f, indeterminate "
        "when 0<|delta|<=epsilon_f, and no modeled benefit when AC is retained or delta>=0. "
        "The global maximum discrepancy and 2 epsilon_f are sensitivity thresholds. These "
        "categories are screening heuristics, not confidence intervals or hypothesis tests."
    )
    document.add_heading("D. Factorial Design and Reproducibility", level=2)
    document.add_paragraph(
        f"The full run contains {int(float(value['scenario_cell_count'])):,} result rows: "
        f"{int(float(value['config_feeders_count']))} feeders, "
        f"{int(float(value['config_topologies_count']))} topologies, "
        f"{int(float(value['config_portfolios_count']))} portfolios, "
        f"{int(float(value['config_native_dc_shares_count']))} native-DC shares, "
        f"{int(float(value['config_spatial_seeds_count']))} spatial allocations, "
        f"{int(float(value['config_converter_efficiencies_count']))} boundary efficiencies, "
        "and five design cells. Randomness is confined to frozen spatial seeds. Raw objects are "
        "preserved separately from processed data and verified by file size and SHA-256. "
        "Every manuscript scalar is written to a machine-readable registry before document "
        "generation."
    )
    document.add_paragraph(
        f"Of {int(float(value['joint_cell_count'])):,} joint cells, "
        f"{int(float(value['joint_repeated_outcome_group_rows']))} rows belong to outcome "
        "groups repeated after excluding the spatial seed, and "
        f"{int(float(value['joint_unique_outcome_rows_excluding_seed']))} distinct outcome "
        "rows remain under that grouping. Accordingly, architecture frequencies summarize "
        "the frozen design grid; repeated structural values are not treated as independent "
        "statistical replications and no design-grid proportion is interpreted as a "
        "real-world probability."
    )
    document.add_heading("III. RESULTS", level=1)
    document.add_heading("A. Attribution Under the Central Setting", level=2)
    document.add_paragraph(
        f"The central cell uses K3, {float(value['central_native_dc_share_pct']):.0f}% "
        f"native-DC annual energy, {float(value['central_boundary_efficiency_pct']):.0f}% "
        f"boundary efficiency, and seed {int(float(value['central_spatial_seed']))}. On the "
        "original topology, all three feeder optima are hybrid. Relative to inherited AC, "
        f"conductor optimization changed mean modeled loss by "
        f"{float(value['mean_conductor_effect_pct']):.2f}%. Relative to matched "
        "conductor-optimized AC, joint domain/conductor planning changed mean modeled loss "
        f"by {float(value['mean_hybrid_vs_ac_loss_change_pct']):.2f}%. These are modeled "
        "attribution contrasts, not claims about realized utility savings."
    )
    document.add_paragraph(
        "Table III reports the six central feeder/topology cells, and Fig. 2 visualizes the "
        "original-topology decomposition."
    )
    attribution_rows = []
    for feeder in ("rural", "suburban", "urban"):
        for topology, topology_label in (("original", "Orig."), ("synthetic_mst", "MST")):
            prefix = f"central_{feeder}_{topology}"
            attribution_rows.append(
                [
                    f"{feeder.title()} {topology_label}",
                    f"{float(value[f'{prefix}_inherited_ac_loss_kwh']):.0f}",
                    f"{float(value[f'{prefix}_conductor_optimized_ac_loss_kwh']):.0f}",
                    f"{float(value[f'{prefix}_fixed_conductor_hybrid_loss_kwh']):.0f}",
                    f"{float(value[f'{prefix}_joint_hybrid_loss_kwh']):.0f}",
                    f"{float(value[f'{prefix}_joint_incremental_architecture_contrast_pct']):.1f}%",
                ]
            )
    _add_table(
        document,
        ["Feeder/topology", "AC base", "AC opt.", "H fixed", "H joint", "Joint contrast"],
        attribution_rows,
    )
    document.add_paragraph(
        "TABLE III. Central annual modeled loss (kWh/year). Joint contrast is relative to "
        "matched conductor-optimized AC.",
        style="Figure Caption",
    )
    document.add_paragraph(
        "The nested cells distinguish three mechanisms. In the rural original feeder, the "
        f"conductor contrast is "
        f"{float(value['central_rural_original_conductor_contrast_pct']):.2f}%, fixed-design "
        f"DC permission is {float(value['central_rural_original_fixed_design_dc_permission_contrast_pct']):.2f}%, "
        f"and the joint architecture contrast is "
        f"{float(value['central_rural_original_joint_incremental_architecture_contrast_pct']):.2f}%. "
        "The corresponding joint contrasts are "
        f"{float(value['central_suburban_original_joint_incremental_architecture_contrast_pct']):.2f}% "
        "for suburban and "
        f"{float(value['central_urban_original_joint_incremental_architecture_contrast_pct']):.2f}% "
        "for urban. Thus conductor optimization and domain permission are neither identical "
        "nor additive in general."
    )
    document.add_paragraph(
        "Table IV reports the central boundary characteristics and the matched topology "
        "sensitivity."
    )
    topology_rows = []
    for feeder in ("rural", "suburban", "urban"):
        original_prefix = f"central_{feeder}_original"
        synthetic_prefix = f"central_{feeder}_synthetic_mst"
        topology_rows.append(
            [
                feeder.title(),
                str(value[f"{original_prefix}_joint_architecture"]),
                f"{float(value[f'{original_prefix}_joint_dc_bus_fraction']) * 100:.1f}%",
                str(int(float(value[f"{original_prefix}_joint_boundary_count"]))),
                f"{float(value[f'{synthetic_prefix}_topology_sensitivity_vs_original_joint_pct']):.1f}%",
            ]
        )
    _add_table(
        document,
        ["Feeder", "Class", "DC buses", "Boundaries", "MST joint-loss change"],
        topology_rows,
    )
    document.add_paragraph(
        "TABLE IV. Central boundary characteristics on the original topology and synthetic-MST "
        "sensitivity. The topology column is a modeled sensitivity, not an observed redesign.",
        style="Figure Caption",
    )
    document.add_paragraph(
        "Synthetic minimum-length topology increased central joint loss relative to original "
        f"topology by {float(value['central_rural_synthetic_mst_topology_sensitivity_vs_original_joint_pct']):.2f}% "
        f"(rural), {float(value['central_suburban_synthetic_mst_topology_sensitivity_vs_original_joint_pct']):.2f}% "
        f"(suburban), and {float(value['central_urban_synthetic_mst_topology_sensitivity_vs_original_joint_pct']):.2f}% "
        "(urban). This result illustrates why a topology change cannot be silently bundled "
        "with DC permission. The table also shows that the central original-topology solutions "
        "use small DC bus fractions and multiple interior boundaries rather than an automatic "
        "feeder-wide conversion."
    )
    document.add_picture(str(FIGURES / "figure_1_decomposition.png"), width=Inches(3.1))
    document.add_paragraph(
        "Fig. 2. Matched decomposition of annual modeled losses under the central setting.",
        style="Figure Caption",
    )
    document.add_heading("B. Conditional Architecture Transitions", level=2)
    document.add_paragraph(
        "The endogenous boundary moved monotonically only in portions of the parameter space. "
        "At low native-DC share, endpoint mismatch and boundary standby losses often retained "
        "all-AC. Higher native-DC share expanded DC subtrees, but topology, feeder geometry, "
        "portfolio, and spatial clustering changed both the transition location and boundary "
        "signature. The same aggregate native-DC share can therefore yield different domain "
        "assignments."
    )
    document.add_paragraph(
        f"Across all joint-optimization cells, a hybrid architecture was selected in "
        f"{float(value['joint_hybrid_selected_pct']):.1f}% of cells, all-AC in "
        f"{float(value['joint_all_ac_selected_pct']):.1f}%, and root-converted DC in "
        f"{float(value['joint_root_dc_selected_pct']):.1f}%. These are frequencies in the "
        "sampled design grid, not probabilities for real feeders. They show that endpoint "
        "architectures remain active and that the optimizer does not mechanically force a "
        "mixed solution."
    )
    document.add_paragraph(
        "Table V separates those counts by boundary efficiency. Figs. 3 and 4 then show a "
        "representative transition surface and the aggregate DC-assignment response."
    )
    _add_table(
        document,
        ["Boundary efficiency", "All-AC", "Hybrid", "Root DC"],
        [
            [
                f"{efficiency}%",
                str(int(float(value[f"joint_efficiency_{efficiency}_all_ac_cells"]))),
                str(int(float(value[f"joint_efficiency_{efficiency}_hybrid_cells"]))),
                str(int(float(value[f"joint_efficiency_{efficiency}_root_dc_cells"]))),
            ]
            for efficiency in (94, 96, 98, 99)
        ],
    )
    document.add_paragraph(
        "TABLE V. Joint architecture counts by boundary-efficiency scenario. Each row contains "
        f"{int(float(value['joint_cells_per_efficiency']))} frozen design cells.",
        style="Figure Caption",
    )
    document.add_paragraph(
        "At 94% and 96% boundary efficiency, every tested joint cell retained all-AC. Hybrid "
        "and root-converted DC solutions appeared at 98% and 99%, but nominal efficiency did "
        "not uniquely determine the boundary because standby loss scales with installed "
        "capacity and endpoint mismatch depends on utilization and location. Transition "
        "locations and signatures for every feeder, topology, portfolio, share, seed, and "
        "efficiency are retained in the machine-readable factorial tables."
    )
    document.add_picture(str(FIGURES / "figure_2_phase_diagram.png"), width=Inches(3.0))
    document.add_paragraph(
        "Fig. 3. Suburban architecture transition under K3; values are bus fractions assigned DC.",
        style="Figure Caption",
    )
    document.add_picture(str(FIGURES / "figure_3_architecture_transition.png"), width=Inches(3.0))
    document.add_paragraph(
        "Fig. 4. Mean DC-assigned bus share versus native-DC energy share at 98% efficiency.",
        style="Figure Caption",
    )
    document.add_paragraph(
        "Table VI stratifies the joint architecture counts by native-DC annual-energy share."
    )
    _add_table(
        document,
        ["Native-DC energy", "All-AC", "Hybrid", "Root DC", "Cells"],
        [
            [
                f"{native_share}%",
                str(int(float(value[f"joint_native_share_{native_share}_all_ac_cells"]))),
                str(int(float(value[f"joint_native_share_{native_share}_hybrid_cells"]))),
                str(int(float(value[f"joint_native_share_{native_share}_root_dc_cells"]))),
                str(int(float(value["joint_cells_per_native_share"]))),
            ]
            for native_share in (0, 25, 50, 75, 100)
        ],
    )
    document.add_paragraph(
        "TABLE VI. Joint architecture counts by native-DC annual-energy share. Counts pool "
        "feeders, topologies, portfolios, spatial seeds, and boundary efficiencies.",
        style="Figure Caption",
    )
    document.add_paragraph(
        "Table VI identifies three conditional regions without assigning probabilities to "
        "them. All cells remain AC when no endpoint energy is native DC. Mixed architectures "
        "appear at intermediate shares, while complete root conversion occurs only in a "
        "subset of the 100% native-DC cells. The exact balance remains converter- and "
        "network-dependent, so native-DC share is a necessary but insufficient boundary rule."
    )
    document.add_heading("C. Exactness and Model-Discrepancy Screening", level=2)
    document.add_paragraph(
        "On independently enumerated toy trees, the dynamic program matched the exhaustive "
        f"minimum objective in all {int(float(value['dp_exactness_case_count']))} declared "
        "cases. Unique optima also matched in architecture and conductor selection. The cases "
        "cover all-AC, root-converted DC, interior hybrid, high converter penalty, high "
        "endpoint mismatch, a conductor-driven architecture change, and an explicit tied "
        f"optimum. The maximum absolute objective difference was "
        f"{float(value['dp_exactness_max_objective_error']):.2e}; "
        f"{int(float(value['dp_exactness_tied_case_count']))} case contained multiple "
        "brute-force optima."
    )
    document.add_paragraph(
        f"All three pandapower cases converged. The median absolute discrepancy between the "
        f"balanced resistive planning approximation and nonlinear AC line loss was "
        f"{float(value['median_validation_error_pct']):.3f}%, with a maximum of "
        f"{float(value['maximum_validation_error_pct']):.3f}%. The feeder values were "
        f"{float(value['validation_rural_relative_error_pct']):.3f}% (rural), "
        f"{float(value['validation_suburban_relative_error_pct']):.3f}% (suburban), and "
        f"{float(value['validation_urban_relative_error_pct']):.3f}% (urban), using nonlinear "
        "pandapower line loss as denominator. Fig. 5 displays the matched nonlinear and "
        "screening values."
    )
    document.add_picture(str(FIGURES / "figure_4_validation.png"), width=Inches(2.9))
    document.add_paragraph(
        "Fig. 5. All-AC approximation versus pandapower at the benchmark operating point.",
        style="Figure Caption",
    )
    document.add_paragraph(
        "Table VII compares the primary discrepancy rule with the two conservative "
        "sensitivity thresholds."
    )
    _add_table(
        document,
        ["Screening rule", "No benefit/AC", "Robust", "Indeterminate"],
        [
            [
                "Feeder epsilon",
                str(int(float(value["discrepancy_feeder_no_benefit_cells"]))),
                str(int(float(value["discrepancy_feeder_robust_cells"]))),
                str(int(float(value["discrepancy_feeder_indeterminate_cells"]))),
            ],
            [
                "Global max epsilon",
                str(int(float(value["discrepancy_global_max_no_benefit_cells"]))),
                str(int(float(value["discrepancy_global_max_robust_cells"]))),
                str(int(float(value["discrepancy_global_max_indeterminate_cells"]))),
            ],
            [
                "Twice feeder epsilon",
                str(int(float(value["discrepancy_double_feeder_no_benefit_cells"]))),
                str(int(float(value["discrepancy_double_feeder_robust_cells"]))),
                str(int(float(value["discrepancy_double_feeder_indeterminate_cells"]))),
            ],
        ],
    )
    document.add_paragraph(
        f"TABLE VII. Classification of {int(float(value['joint_cell_count'])):,} joint cells "
        "under alternative approximation-"
        "discrepancy screens.",
        style="Figure Caption",
    )
    document.add_paragraph(
        f"Under the primary feeder-specific rule, "
        f"{int(float(value['discrepancy_feeder_robust_cells']))} cells were robustly material, "
        f"{int(float(value['discrepancy_feeder_indeterminate_cells']))} were indeterminate, "
        f"and {int(float(value['discrepancy_feeder_no_benefit_cells']))} retained AC or had "
        "no modeled loss benefit. Doubling the feeder discrepancy moved cases from robust to "
        "indeterminate without changing the no-benefit count. The screen therefore prevents "
        "small negative contrasts from being presented as resolved architecture improvements."
    )
    document.add_paragraph(
        "The qualitative conclusion is stable to the two more conservative screening "
        "thresholds, although the robust/indeterminate boundary moves as intended."
    )
    document.add_paragraph(
        "Table VIII stratifies the primary discrepancy classification by benchmark feeder."
    )
    _add_table(
        document,
        ["Feeder", "No benefit/AC", "Robust", "Indeterminate"],
        [
            [
                feeder.title(),
                str(int(float(value[f"discrepancy_feeder_rule_{feeder}_no_benefit_cells"]))),
                str(int(float(value[f"discrepancy_feeder_rule_{feeder}_robust_cells"]))),
                str(int(float(value[f"discrepancy_feeder_rule_{feeder}_indeterminate_cells"]))),
            ]
            for feeder in ("rural", "suburban", "urban")
        ],
    )
    document.add_paragraph(
        "TABLE VIII. Feeder-stratified primary discrepancy classification; each feeder "
        f"contributes {int(float(value['joint_cells_per_feeder']))} joint cells.",
        style="Figure Caption",
    )
    document.add_paragraph(
        "The feeder stratification separates approximation screening from feeder prevalence. "
        "It is useful for "
        "checking whether the aggregate count is dominated by one benchmark, but it remains a "
        "balanced scenario-grid summary rather than a statistical estimate for rural, "
        "suburban, or urban feeder populations."
    )
    document.add_heading("D. Converter-Component Ablations", level=2)
    document.add_paragraph(
        "Eight matched ablations remove individual conversion terms while preserving the same "
        "factorial cells. The comparisons isolate boundary-efficiency loss, standby loss, the "
        "converter-capacity regularizer, endpoint mismatch, all boundary electrical loss, and "
        "two idealized upper bounds. The assumed ranges are dimensionless planning scenarios; "
        "they are not manufacturer curves or equipment-cost estimates."
    )
    document.add_paragraph(
        "Table IX summarizes architecture, boundary, DC-assignment, and loss changes relative "
        "to the matched baseline."
    )
    _add_table(
        document,
        ["Ablation", "Architecture change", "Boundary change", "Mean DC buses", "Mean loss contrast"],
        [
            [
                "Baseline",
                f"{float(value['ablation_baseline_architecture_change_pct']):.1f}%",
                f"{float(value['ablation_baseline_boundary_change_pct']):.1f}%",
                f"{float(value['ablation_baseline_mean_dc_bus_fraction']) * 100:.1f}%",
                f"{float(value['ablation_baseline_mean_loss_change_pct']):.1f}%",
            ],
            [
                "No capacity reg.",
                f"{float(value['ablation_no_capacity_regularizer_architecture_change_pct']):.1f}%",
                f"{float(value['ablation_no_capacity_regularizer_boundary_change_pct']):.1f}%",
                f"{float(value['ablation_no_capacity_regularizer_mean_dc_bus_fraction']) * 100:.1f}%",
                f"{float(value['ablation_no_capacity_regularizer_mean_loss_change_pct']):.1f}%",
            ],
            [
                "No standby",
                f"{float(value['ablation_no_standby_loss_architecture_change_pct']):.1f}%",
                f"{float(value['ablation_no_standby_loss_boundary_change_pct']):.1f}%",
                f"{float(value['ablation_no_standby_loss_mean_dc_bus_fraction']) * 100:.1f}%",
                f"{float(value['ablation_no_standby_loss_mean_loss_change_pct']):.1f}%",
            ],
            [
                "No boundary eta loss",
                f"{float(value['ablation_no_boundary_efficiency_loss_architecture_change_pct']):.1f}%",
                f"{float(value['ablation_no_boundary_efficiency_loss_boundary_change_pct']):.1f}%",
                f"{float(value['ablation_no_boundary_efficiency_loss_mean_dc_bus_fraction']) * 100:.1f}%",
                f"{float(value['ablation_no_boundary_efficiency_loss_mean_loss_change_pct']):.1f}%",
            ],
            [
                "Electrical ideal",
                f"{float(value['ablation_electrical_ideal_conversion_architecture_change_pct']):.1f}%",
                f"{float(value['ablation_electrical_ideal_conversion_boundary_change_pct']):.1f}%",
                f"{float(value['ablation_electrical_ideal_conversion_mean_dc_bus_fraction']) * 100:.1f}%",
                f"{float(value['ablation_electrical_ideal_conversion_mean_loss_change_pct']):.1f}%",
            ],
            [
                "Full ideal converter",
                f"{float(value['ablation_full_ideal_converter_architecture_change_pct']):.1f}%",
                f"{float(value['ablation_full_ideal_converter_boundary_change_pct']):.1f}%",
                f"{float(value['ablation_full_ideal_converter_mean_dc_bus_fraction']) * 100:.1f}%",
                f"{float(value['ablation_full_ideal_converter_mean_loss_change_pct']):.1f}%",
            ],
        ],
    )
    document.add_paragraph(
        "TABLE IX. Aggregate converter ablations across the matched factorial grid.",
        style="Figure Caption",
    )
    document.add_paragraph(
        "The capacity regularizer changed architecture in only "
        f"{float(value['ablation_no_capacity_regularizer_architecture_change_pct']):.2f}% of "
        "cells, although more boundary signatures changed. Removing standby loss changed "
        f"architecture in {float(value['ablation_no_standby_loss_architecture_change_pct']):.2f}% "
        f"and boundaries in {float(value['ablation_no_standby_loss_boundary_change_pct']):.2f}%. "
        "Removing boundary-efficiency loss produced much larger changes, while eliminating "
        "endpoint mismatch selected all-AC in "
        f"{float(value['ablation_no_endpoint_mismatch_loss_all_ac_pct']):.0f}% of cells. The "
        "last result is mechanistic: without avoided endpoint conversion, the modeled reason "
        "to assign loads to DC disappears under the remaining penalties."
    )
    document.add_paragraph(
        "The electrical-ideal and full-ideal cases should not be read as feasible forecasts. "
        f"The electrical-ideal case changed architecture in "
        f"{float(value['ablation_electrical_ideal_conversion_architecture_change_pct']):.1f}% "
        f"of cells; the full ideal case changed it in "
        f"{float(value['ablation_full_ideal_converter_architecture_change_pct']):.1f}% and "
        f"selected root-converted DC in "
        f"{float(value['ablation_full_ideal_converter_root_dc_pct']):.1f}%. Their purpose is "
        "to bound conversion penalties and expose utilization dependence."
    )
    document.add_paragraph(
        "The ablation table therefore supports a component-level interpretation rather than a single "
        "undifferentiated converter-sensitivity claim."
    )
    document.add_heading("E. Performance and Nonlinear Stress Checks", level=2)
    document.add_paragraph(
        f"The implementation ran on {value['performance_platform']} with Python "
        f"{value['performance_python']} and {value['performance_processor']}; "
        f"{int(float(value['performance_logical_cpus']))} logical CPUs were visible. Software "
        f"versions were {value['performance_software_versions']}. Peak process memory was "
        f"{float(value['performance_peak_process_rss_mb']):.0f} MiB. The "
        f"{int(float(value['scenario_cell_count'])):,}-row factorial required "
        f"{float(value['performance_factorial_scenario_runtime_s']):.1f} s, and the "
        "converter-ablation run required "
        f"{float(value['performance_converter_ablation_runtime_s']):.1f} s in this environment."
    )
    document.add_paragraph(
        "Table X links the analytical state counts to measured solve times for the original "
        "benchmark topologies."
    )
    _add_table(
        document,
        ["Feeder", "Nodes/edges", "Median solve", "DP states", "Transitions"],
        [
            [
                feeder.title(),
                f"{int(float(value[f'performance_{feeder}_original_nodes']))}/"
                f"{int(float(value[f'performance_{feeder}_original_edges']))}",
                f"{float(value[f'performance_{feeder}_original_median_runtime_ms']):.2f} ms",
                str(int(float(value[f"performance_{feeder}_original_state_evaluations"]))),
                str(int(float(value[f"performance_{feeder}_original_transition_evaluations"]))),
            ]
            for feeder in ("rural", "suburban", "urban")
        ],
    )
    document.add_paragraph(
        "TABLE X. Original-topology central solve evidence; medians use "
        f"{int(float(value['performance_benchmark_repeats']))} repeated solves.",
        style="Figure Caption",
    )
    document.add_paragraph(
        "The analytical transition count is O(|E||D|^2|S|), with two domain states and three "
        "conductor scales in the present configuration. Median repeated-solve time across all "
        f"feeder/topology cells ranged from "
        f"{float(value['performance_min_feeder_median_runtime_ms']):.2f} to "
        f"{float(value['performance_max_feeder_median_runtime_ms']):.2f} ms. On chain toys, "
        f"exhaustive enumeration reached {int(float(value['performance_max_enumerated_designs'])):,} "
        f"designs while the maximum DP-versus-enumeration objective difference remained "
        f"{float(value['performance_toy_max_objective_error']):.2e}. Timing is implementation "
        "evidence, not an independent proof of asymptotic complexity."
    )
    document.add_paragraph(
        f"All {int(float(value['validation_stress_case_count']))} balanced all-AC feeder/"
        "portfolio stress cases converged. Across the grid, minimum voltage was "
        f"{float(value['validation_stress_min_voltage_pu']):.3f} p.u., maximum line loading "
        f"was {float(value['validation_stress_max_loading_pct']):.1f}%, and maximum line-loss "
        f"discrepancy was {float(value['validation_stress_max_error_pct']):.3f}%. These checks "
        "support the all-AC approximation over the configured peak multipliers; they do not "
        "validate hybrid converter controls, unbalance, harmonics, grounding, protection, or "
        "reliability."
    )
    document.add_heading("F. Distributional and Feasibility Diagnostics", level=2)
    document.add_paragraph(
        f"Across all paired realistic cells, the median joint-hybrid change relative to "
        f"conductor-optimized AC was {float(value['paired_loss_change_median_pct']):.2f}%; "
        f"the 5th–95th percentile interval was "
        f"{float(value['paired_loss_change_q05_pct']):.2f}% to "
        f"{float(value['paired_loss_change_q95_pct']):.2f}%. The interval spans materially "
        "different outcomes because native-DC share changes endpoint mismatch and converter "
        "utilization. Portfolio peak factors shift line losses quadratically, whereas annual "
        "energy factors shift endpoint and conversion losses approximately linearly."
    )
    document.add_paragraph(
        f"The screening diagnostics flagged {int(float(value['constraint_violation_cells']))} "
        "joint-optimization cells above the configured ampacity-utilization or single-edge "
        "voltage-drop limit. The limits were "
        f"{float(value['config_ampacity_limit_fraction']):.2f} p.u. and "
        f"{float(value['config_voltage_drop_limit_pu']):.2f} p.u. Flagged cells remain in the "
        "public results so that "
        "an apparent loss benefit is not presented as a feasible design. A utility-grade study "
        "would enforce phase-specific voltages, thermal ratings, and contingency requirements "
        "inside the optimization."
    )
    document.add_heading("IV. DISCUSSION", level=1)
    document.add_heading("A. What the Attribution Design Changes", level=2)
    document.add_paragraph(
        "The principal finding is methodological: the value assigned to DC depends on the "
        "comparison cell. Conductor redesign can move losses before any bus changes domain, "
        "while a topology sensitivity changes downstream power and line length. Reporting only "
        "the inherited-AC versus joint-hybrid difference would combine these mechanisms. The "
        "matched design prevents that total from being labeled a DC effect."
    )
    document.add_paragraph(
        "This distinction matters for both retrofit and greenfield interpretation. In a "
        "retrofit, fixed-design DC permission asks whether domain reassignment has marginal "
        "planning value without assuming simultaneous conductor replacement. In a constrained "
        "redesign, the joint contrast asks what is gained after AC receives the same conductor "
        "menu. In a synthetic greenfield sensitivity, topology is varied explicitly and its "
        "change is reported separately. None of these estimands is a causal treatment effect; "
        "each is a paired contrast under frozen model assumptions."
    )
    document.add_heading("B. Why Converter Utilization Matters", level=2)
    document.add_paragraph(
        "The transition maps also reject a universal boundary rule. Native-DC penetration "
        "favors DC service by avoiding endpoint conversion, whereas converter standby loss "
        "penalizes small or lightly loaded DC subtrees. A high nominal efficiency is therefore "
        "insufficient evidence; utilization and spatial clustering matter. All-AC remains a "
        "valid optimum in parts of the tested space, and the model preserves that negative "
        "result."
    )
    document.add_paragraph(
        "The selected boundary is therefore a marginal architecture signal. A DC subtree "
        "becomes attractive when avoided endpoint mismatch and lower conductor loss exceed "
        "efficiency, standby, capacity, and material consequences at its boundary. Spatial "
        "clustering is decisive: the same aggregate DC-native share can produce different "
        "boundaries when concentrated behind one branch or dispersed among branches. This "
        "mechanism explains why one system-wide penetration threshold is not transportable "
        "across feeder geometries."
    )
    document.add_paragraph(
        "The ablations sharpen this mechanism. Removing the capacity regularizer rarely "
        "changes the architecture, while removing standby or boundary-efficiency loss changes "
        "many more boundary signatures. Consequently, quoting only a high full-load converter "
        "efficiency is inadequate. A planner also needs the power transferred through the "
        "boundary, the size implied by peak demand, the energy served behind it, and the "
        "spatial concentration of native-DC endpoints. Those variables determine whether "
        "fixed and throughput-dependent conversion penalties are diluted by useful service."
    )
    document.add_heading("C. Interpretation of Robust and Null Regions", level=2)
    document.add_paragraph(
        "The discrepancy classification is intentionally asymmetric. Large modeled benefits "
        "that exceed the all-AC approximation error are labeled robustly material only under "
        "the declared screening rule; smaller benefits remain indeterminate. All-AC and "
        "nonnegative contrasts are retained as no-benefit outcomes. This prevents the "
        "factorial size from turning numerically small differences into persuasive evidence. "
        "It also preserves the scientifically relevant negative result that realistic "
        "conversion accounting can leave AC unchanged over broad portions of the design grid."
    )
    document.add_paragraph(
        "Normalized objective weights also separate engineering screening from economic "
        "valuation. Public cost evidence for low-voltage bidirectional converters, DC "
        "switchgear, protection, and maintenance is too heterogeneous for a defensible "
        "absolute net-present-value estimate. Reporting electrical components and capacity "
        "requirements separately permits later studies to apply region- and vendor-specific "
        "prices without changing the architecture-attribution design."
    )
    document.add_heading("D. Limitations and External Validity", level=2)
    document.add_paragraph(
        "The model is deliberately narrower than a deployable planning tool. It uses balanced "
        "radial active-power accounting and proxy portfolios rather than unbalanced optimal "
        "power flow, phase switching, explicit inverter reactive control, chronological "
        "battery dispatch, harmonics, grounding, fault interruption, protection coordination, "
        "reliability, or regulation. The synthetic spanning tree is not observed geography. "
        "Public evidence for bidirectional LV boundary-converter cost, part-load efficiency, "
        "standby power, and lifetime remains insufficient for absolute lifecycle-cost claims. "
        "Accordingly, this paper reports conditional normalized evidence."
    )
    document.add_paragraph(
        "The nonlinear comparison validates only the all-AC loss approximation. It does not "
        "validate hybrid converter voltage control, bidirectional dispatch, fault behavior, "
        "or interactions among multiple converters. The ampacity and voltage-drop quantities "
        "are post-optimization diagnostics rather than binding constraints. Moreover, the "
        "portfolio multipliers are controlled stress proxies, not probabilistic forecasts, "
        "and the three frozen spatial allocations are sensitivity cases rather than a random "
        "sample of future adoption."
    )
    document.add_paragraph(
        "Benchmark feeders improve reproducibility but limit external validity. SimBench "
        "archetypes and the synthetic minimum-length trees do not represent the distribution "
        "of utility assets, local grounding practice, service conductor arrangements, or "
        "regulatory requirements. The conductor-count convention and the "
        f"{float(value['config_dc_voltage_pole_to_pole_v']):.0f}-V DC scenario are controlled "
        "modeling choices. The reported boundary locations should therefore "
        "be interpreted as mechanism examples within these networks, not transferable siting "
        "recommendations."
    )
    document.add_paragraph(
        "Future work should embed the attribution design in unbalanced multiperiod AC/DC "
        "optimal power flow, use measured bidirectional converter maps, and co-design "
        "protection. Chronological operation should include power direction, storage dispatch, "
        "reactive support, and part-load losses. Field utility data and equipment quotations "
        "would be required before claiming deployable boundary locations, lifecycle economics, "
        "or external validity beyond benchmark mechanism identification."
    )
    document.add_heading("V. CONCLUSION", level=1)
    document.add_paragraph(
        "An endogenous AC/DC boundary is scientifically informative only when the comparison "
        "separates electrical-domain permission from conductor and topology redesign. The "
        "nested contrasts show that these freedoms produce distinct changes, while the full "
        "factorial shows conditional all-AC, hybrid, and root-converted DC regions rather than "
        "a general DC advantage. Discrepancy screening preserves indeterminate and null "
        "outcomes, and converter ablations show why nominal efficiency must be interpreted "
        "together with utilization, standby, and endpoint mismatch. The exact dynamic program "
        "and registry-backed open pipeline provide reproducible mechanism evidence for the "
        "declared balanced radial approximation. Unbalanced hybrid operation, protection, "
        "reliability, measured converter maps, and location-specific economics remain necessary "
        "before any deployment decision."
    )
    document.add_heading("ACKNOWLEDGMENT", level=1)
    document.add_paragraph(
        "[FUNDING, CONFLICTS, AND NONAUTHOR CONTRIBUTIONS TO BE CONFIRMED.] During preparation "
        "of this work, the authors used Cognition Devin for substantive assistance in "
        "generating and debugging analysis code, drafting and editing all manuscript sections, "
        "and preparing figures and submission files. The authors reviewed and verified the "
        "analysis, numerical results, citations, claims, and final text and remain accountable "
        "for the article."
    )
    document.add_heading("DATA AND CODE AVAILABILITY", level=1)
    document.add_paragraph(
        f"The analysis code, configurations, provenance ledger, and generated outputs are "
        f"available at {PUBLIC_REPOSITORY_URL}. SimBench is acquired under ODbL/BSD terms; "
        "the charging data are CC BY 4.0. Exact source URLs, versions, file sizes, and "
        "SHA-256 values are recorded in the machine-readable ledger."
    )
    document.add_heading("REFERENCES", level=1)
    for number, reference in enumerate(references, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.12)
        paragraph.paragraph_format.first_line_indent = Inches(-0.12)
        paragraph.add_run(_reference_text(reference, number))
    _move_table_captions_above_tables(document)
    path = build / "acdc_boundary_tsg.docx"
    document.save(path)
    _canonicalize_office_archive(path)
    return path


def build_cover_letter(values: pd.DataFrame) -> Path:
    build = SUBMISSION / "build"
    build.mkdir(parents=True, exist_ok=True)
    value = _headline_values(values)
    document = Document()
    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(11)
    document.add_paragraph("[DATE]")
    document.add_paragraph("Editor-in-Chief\nIEEE Transactions on Smart Grid")
    document.add_paragraph("Re: Submission of “Attributing AC/DC Boundary Value in Active Low-Voltage Distribution Planning”")
    document.add_paragraph("Dear Editor-in-Chief:")
    document.add_paragraph(
        "Please consider the enclosed manuscript for publication in IEEE Transactions on "
        "Smart Grid. The study asks when a modern low-voltage feeder should introduce a DC "
        "subtree after controlling for topology and conductor redesign. Its contribution is "
        "a matched architecture-attribution framework rather than a claim that AC/DC state "
        "optimization is new."
    )
    document.add_paragraph(
        f"Across three public SimBench feeders, the central realistic-converter comparison "
        f"changed modeled annual loss by {float(value['mean_hybrid_vs_ac_loss_change_pct']):.2f}% "
        "relative to conductor-optimized AC, while preserving all-AC whenever it remained "
        "preferred. The paper emphasizes active-distribution relevance through DER/EV stress "
        "portfolios and conditional architecture transitions. It explicitly limits claims "
        "concerning protection, reliability, observed infrastructure, and absolute economics."
    )
    document.add_paragraph(
        "The complete analysis is configuration-driven. Public source objects are retained "
        "with checksums and licenses, and every manuscript number maps to a machine-readable "
        "registry. The work is original and is not under consideration elsewhere, subject to "
        "confirmation by all authors."
    )
    document.add_paragraph(
        "[AUTHORSHIP, PRIOR PUBLICATION, FUNDING, CONFLICT-OF-INTEREST, ORCID, AND CONTACT "
        "DECLARATIONS MUST BE COMPLETED BY THE AUTHORS.]"
    )
    document.add_paragraph("Sincerely,\n[CORRESPONDING AUTHOR]")
    path = build / "cover_letter.docx"
    document.save(path)
    _canonicalize_office_archive(path)
    return path


def build_submission_documents(values: pd.DataFrame) -> None:
    build = SUBMISSION / "build"
    build.mkdir(parents=True, exist_ok=True)
    checklist = """# IEEE TSG submission checklist

- [ ] Confirm all author names, IEEE membership grades, affiliations, emails, and ORCIDs.
- [ ] Confirm corresponding author and complete title-page footnote.
- [ ] Confirm funding, conflicts of interest, acknowledgments, and prior-publication status.
- [ ] Confirm every author approves the final manuscript.
- [x] Initial manuscript is prepared in US-Letter, double-column format.
- [x] Abstract is 150–200 words and index terms are alphabetized.
- [x] Figures are available as PNG and 600-dpi LZW TIFF.
- [x] References are numbered by first appearance and DOI metadata were checked.
- [x] No conventional supplementary file is included.
- [ ] Recheck the live TSG portal and May 2026 PES AI-disclosure fields immediately before submission.
- [ ] Decide whether to submit an optional graphical abstract.
- [x] Public repository released: https://github.com/bougtoir/acdc_boundary_grid.
- [ ] Archive a version and insert its DOI if required before submission.
- [ ] Human authors must verify all numerical and scientific claims.
"""
    (build / "submission_checklist.md").write_text(checklist, encoding="utf-8")
    declarations = """# Declarations requiring author confirmation

## Authors and affiliations
[TO BE COMPLETED]

## ORCID
[REQUIRED FOR EACH AUTHOR; TO BE COMPLETED]

## Funding
[TO BE COMPLETED; do not infer]

## Conflicts of interest
[TO BE COMPLETED; do not infer]

## Author contributions
[TO BE COMPLETED BY THE AUTHORS]

## Prior publication and concurrent submission
[TO BE CONFIRMED]

## AI assistance
Cognition Devin generated and debugged code, drafted text, and prepared figures and files.
The final Acknowledgment identifies the system, affected content, and substantive level of
use. The authors must verify that wording against the live IEEE PES policy.
"""
    (build / "author_declarations.md").write_text(declarations, encoding="utf-8")
    value_map = _headline_values(values)
    git_result = subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )
    git_commit = git_result.stdout.strip() if git_result.returncode == 0 else "unavailable"
    handoff = f"""FINAL HANDOFF — REVISED TSG PACKAGE

Target journal: IEEE Transactions on Smart Grid
Fallbacks: IEEE Transactions on Power Delivery; International Journal of Electrical Power & Energy Systems
Final title: Attributing AC/DC Boundary Value in Active Low-Voltage Distribution Planning
Build-source Git commit: {git_commit}
Acceptance status: Technical gates A–S PASS; Gate T HUMAN METADATA PENDING.

Genuine contribution:
A matched architecture-attribution framework that estimates the marginal value of
permitting DC after separating conductor and topology redesign, with an endogenous
AC/DC boundary and all-AC retained as a feasible outcome. The contribution is conditional
mechanism attribution, not universal DC superiority or a deployment recommendation.

Datasets:
- SimBench v1.6.1 low-voltage rural1, semiurb4, and urban6 feeders and profiles.
- Figshare EV charging transactions, DOI 10.6084/m9.figshare.22495141.v1.
- Frozen raw objects, checksums, licenses, and retrieval records are in data_raw/ and
  provenance/.

Formulation:
Exact dynamic programming on a balanced radial separable AC/DC-domain approximation with
explicit endpoint and boundary conversion, conductor selection, ampacity, and voltage-drop
diagnostics. All-AC is retained as a feasible endpoint.

Exactness:
- {int(float(value_map['dp_exactness_case_count']))} independently enumerated toy cases.
- Maximum absolute DP/brute-force objective error:
  {float(value_map['dp_exactness_max_objective_error']):.3e}.

Central matched original-topology results
(inherited AC / conductor-optimized AC / fixed-conductor hybrid /
joint hybrid / electrical-ideal hybrid; kWh/year):
- Rural: {float(value_map['central_rural_original_inherited_ac_loss_kwh']):.1f} /
  {float(value_map['central_rural_original_conductor_optimized_ac_loss_kwh']):.1f} /
  {float(value_map['central_rural_original_fixed_conductor_hybrid_loss_kwh']):.1f} /
  {float(value_map['central_rural_original_joint_hybrid_loss_kwh']):.1f} /
  {float(value_map['central_rural_original_electrical_ideal_hybrid_loss_kwh']):.1f}.
- Suburban: {float(value_map['central_suburban_original_inherited_ac_loss_kwh']):.1f} /
  {float(value_map['central_suburban_original_conductor_optimized_ac_loss_kwh']):.1f} /
  {float(value_map['central_suburban_original_fixed_conductor_hybrid_loss_kwh']):.1f} /
  {float(value_map['central_suburban_original_joint_hybrid_loss_kwh']):.1f} /
  {float(value_map['central_suburban_original_electrical_ideal_hybrid_loss_kwh']):.1f}.
- Urban: {float(value_map['central_urban_original_inherited_ac_loss_kwh']):.1f} /
  {float(value_map['central_urban_original_conductor_optimized_ac_loss_kwh']):.1f} /
  {float(value_map['central_urban_original_fixed_conductor_hybrid_loss_kwh']):.1f} /
  {float(value_map['central_urban_original_joint_hybrid_loss_kwh']):.1f} /
  {float(value_map['central_urban_original_electrical_ideal_hybrid_loss_kwh']):.1f}.
- Mean conductor redesign contrast: {float(value_map['mean_conductor_effect_pct']):.4f}%.
- Mean joint-hybrid versus conductor-optimized AC contrast:
  {float(value_map['mean_hybrid_vs_ac_loss_change_pct']):.4f}%.
- Central architectures: {value_map['central_architecture_summary']}.

Full-factorial evidence:
- {int(float(value_map['scenario_cell_count']))} scenario rows and
  {int(float(value_map['joint_cell_count']))} joint-optimization cells.
- All-AC / hybrid / root-converted DC:
  {int(float(value_map['joint_all_ac_selected_cells']))} /
  {int(float(value_map['joint_hybrid_selected_cells']))} /
  {int(float(value_map['joint_root_dc_selected_cells']))} cells.
- Feeder-specific discrepancy screening:
  {int(float(value_map['discrepancy_feeder_no_benefit_cells']))} no modeled benefit /
  {int(float(value_map['discrepancy_feeder_robust_cells']))} robustly material /
  {int(float(value_map['discrepancy_feeder_indeterminate_cells']))} indeterminate cells.

Converter mechanisms:
- Removing standby loss changes architecture in
  {float(value_map['ablation_no_standby_loss_architecture_change_pct']):.3f}% of cells.
- Removing boundary-efficiency loss changes architecture in
  {float(value_map['ablation_no_boundary_efficiency_loss_architecture_change_pct']):.3f}%.
- Electrical-ideal conversion changes architecture in
  {float(value_map['ablation_electrical_ideal_conversion_architecture_change_pct']):.3f}%.
- Full ideal conversion changes architecture in
  {float(value_map['ablation_full_ideal_converter_architecture_change_pct']):.3f}%.

Validation and computation:
- Median benchmark pandapower line-loss discrepancy:
  {float(value_map['median_validation_error_pct']):.4f}%.
- {int(float(value_map['validation_stress_case_count']))} all-AC stress cases; maximum
  discrepancy {float(value_map['validation_stress_max_error_pct']):.4f}%.
- Factorial runtime {float(value_map['performance_factorial_scenario_runtime_s']):.2f} s;
  ablation runtime {float(value_map['performance_converter_ablation_runtime_s']):.2f} s.
- Runtime and memory values are live implementation measurements, not complexity proofs.

Limitations:
Balanced active-power approximation; proxy DER portfolios; no protection, harmonics,
grounding, reliability, or deployable absolute cost claims. Synthetic MST topologies are
not observed infrastructure. Design-grid frequencies are not deployment probabilities.
Normalized regularizers are not monetary costs. Planning differences below model discrepancy
are indeterminate. Nonlinear validation is all-AC only and does not validate hybrid controls.

Reproduction:
make setup
make acquire
make all
make test
make qc

Public repository:
https://github.com/bougtoir/acdc_boundary_grid

Final inventory:
- manuscript/build/acdc_boundary_tsg.docx
- manuscript/build/pdf/acdc_boundary_tsg.pdf
- submission/build/cover_letter.docx
- submission/build/author_declarations.md
- submission/build/submission_checklist.md
- submission/build/FINAL_HANDOFF_REVISED.txt
- submission/build/tsg_submission_package.zip
- figures/*.png and figures/*.tiff
- tables/*.csv
- results/manuscript_values.csv and results/result_manifest.json
- qc/*_revision.md, qc/final_acceptance_gates.md, and handoffs/PHASE_R*_HANDOFF.md

Human input required before submission:
Author names, affiliations, emails, ORCIDs, corresponding author, funding, conflicts,
author contributions, prior-publication status, archival DOI if required, and final
verification of IEEE PES AI-disclosure wording.
"""
    (build / "FINAL_HANDOFF.txt").write_text(handoff, encoding="utf-8")
    (build / "FINAL_HANDOFF_REVISED.txt").write_text(handoff, encoding="utf-8")


def docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    return re.sub("<[^>]+>", " ", xml)


def write_result_manifest() -> None:
    records = []
    for path in sorted(RESULTS.glob("*")):
        if path.is_file():
            records.append({"file": path.name, "bytes": path.stat().st_size})
    (RESULTS / "result_manifest.json").write_text(
        json.dumps({"files": records}, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
