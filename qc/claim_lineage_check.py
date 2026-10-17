#!/usr/bin/env python3
"""Phase 3: claim-to-source ledger.

Extracts every numeric token from the rendered manuscript DOCX text and checks
whether it is (a) a registry value from results/manuscript_values.csv,
(b) a declared config/model constant, or (c) unexplained — i.e. a candidate
hardcoded result. Writes qc/claim_to_source_ledger.csv.
"""
import csv
import re
import sys
from pathlib import Path

import yaml
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "manuscript/build/acdc_boundary_tsg.docx"
REG = ROOT / "results/manuscript_values.csv"
CFG = ROOT / "configs/analysis_plan.yaml"
OUT = ROOT / "qc/claim_to_source_ledger.csv"

TOKEN = re.compile(r"(?<![\w/.-])(-?\d[\d,]*(?:\.\d+)?(?:e[+-]?\d+)?)(?![\w/.-])", re.I)

def registry_values():
    vals = {}
    with REG.open() as h:
        for row in csv.DictReader(h):
            try:
                v = float(row["value"])
            except (KeyError, ValueError):
                continue
            for rendered in {round(v, 4), round(v, 3), round(v, 2), round(v, 1), round(v),
                             int(round(v)),
                             round(v * 100, 1), round(v * 100, 2), int(round(v * 100)),
                             float(f"{v:.3g}"), float(f"{v:.4g}")}:
                vals.setdefault(float(rendered), []).append(row["value_id"])
    return vals

def config_literals():
    text = CFG.read_text()
    out = {}
    for m in TOKEN.finditer(text):
        try:
            v = float(m.group(1).replace(",", ""))
        except ValueError:
            continue
        for rendered in {v, v * 100, round(v * 100), v / 2}:
            out.setdefault(float(rendered), []).append("configs/analysis_plan.yaml")
    return out

BENIGN = {0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 100, 1000, 8760, 2026, 12, 15, 18, 24, 25}
BENIGN_NOTE = "enumeration/format/model-defining constant"

def main():
    doc = Document(DOCX)
    reg = registry_values()
    cfg = config_literals()
    rows = []
    texts = [(p.text, p.style.name if p.style else "") for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append((cell.text, "TableCell"))
    for text, style in texts:
        for m in TOKEN.finditer(text):
            lit = float(m.group(1).replace(",", ""))
            ctx = text[max(0, m.start() - 50):m.end() + 50].strip()
            if style in ("Bibliography", "References") or "doi:" in ctx.lower() or "Art. no." in ctx:
                source, verdict = "reference entry", "bibliography"
            elif lit in BENIGN:
                source, verdict = "format/model constant", BENIGN_NOTE
            elif lit in reg:
                source, verdict = "; ".join(sorted(set(reg[lit]))[:3]), "registry"
            elif lit in cfg:
                source, verdict = "configs/analysis_plan.yaml", "config"
            else:
                source, verdict = "", "UNMATCHED"
            rows.append((style, lit, ctx[:110], source, verdict))
    with OUT.open("w", newline="") as h:
        w = csv.writer(h)
        w.writerow(["style", "literal", "context", "source_keys", "verdict"])
        w.writerows(rows)
    unmatched = [r for r in rows if r[4] == "UNMATCHED"]
    print(f"total tokens {len(rows)}; unmatched {len(unmatched)}")
    for r in unmatched[:80]:
        print("UNMATCHED:", r[1], "|", r[2])

if __name__ == "__main__":
    main()
